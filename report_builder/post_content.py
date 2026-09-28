"""
report_builder/post_content.py
References and Appendix for SentixAI Academic Report.
"""

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from report_builder.common import (
    add_chapter_title, add_section_h1, add_section_h2,
    add_body_p, add_bullet_item, add_code_block
)

def build_post_content(doc):
    # ---------------------------------------------------------
    # REFERENCES
    # ---------------------------------------------------------
    doc.add_page_break()
    p_ref = doc.add_paragraph()
    p_ref.alignment = doc.add_paragraph().alignment
    p_ref.paragraph_format.space_before = Pt(12)
    p_ref.paragraph_format.space_after = Pt(18)
    run_ref = p_ref.add_run("REFERENCES")
    run_ref.font.name = 'Times New Roman'
    run_ref.font.size = Pt(16)
    run_ref.bold = True

    references = [
        "[1] B. Liu, \"Sentiment Analysis and Opinion Mining,\" Synthesis Lectures on Human Language Technologies, vol. 5, no. 1, pp. 1-167, Morgan & Claypool Publishers, 2012.",
        "[2] S. L. Loria, \"TextBlob: Simplified Text Processing,\" Secondary TextBlob: Simplified Text Processing, 2018. [Online]. Available: https://textblob.readthedocs.io/.",
        "[3] S. Bird, E. Klein, and E. Loper, \"Natural Language Processing with Python: Analyzing Text with the Natural Language Toolkit,\" O'Reilly Media, Sebastopol, CA, 2009.",
        "[4] M. Grinberg, \"Flask Web Development: Developing Web Applications with Python,\" 2nd ed., O'Reilly Media, Sebastopol, CA, 2018.",
        "[5] P. DuBois, \"MySQL: Developer's Library,\" 5th ed., Addison-Wesley Professional, Boston, MA, 2013.",
        "[6] A. Pak and P. Paroubek, \"Twitter as a Corpus for Sentiment Analysis and Opinion Mining,\" in Proceedings of the Seventh International Conference on Language Resources and Evaluation (LREC'10), Valletta, Malta, 2010.",
        "[7] E. Cambria, D. Das, S. Bandyopadhyay, and A. Feraco, \"A Practical Guide to Sentiment Analysis,\" Socio-Affective Computing, vol. 5, Springer International Publishing, 2017.",
        "[8] R. S. Pressman and B. R. Maxim, \"Software Engineering: A Practitioner's Approach,\" 9th ed., McGraw-Hill Education, New York, NY, 2020.",
        "[9] PythonAnywhere LLP, \"Deploying Flask Applications on PythonAnywhere,\" PythonAnywhere Documentation, 2024. [Online]. Available: https://help.pythonanywhere.com/pages/Flask/.",
        "[10] Chart.js Community, \"Chart.js: Open source HTML5 Charts for most popular charting features,\" 2024. [Online]. Available: https://www.chartjs.org/docs/."
    ]

    for ref in references:
        add_body_p(doc, ref)

    # ---------------------------------------------------------
    # APPENDIX
    # ---------------------------------------------------------
    doc.add_page_break()
    p_app = doc.add_paragraph()
    p_app.paragraph_format.space_before = Pt(12)
    p_app.paragraph_format.space_after = Pt(18)
    run_app = p_app.add_run("APPENDIX: USER MANUAL & DEPLOYMENT REFERENCE")
    run_app.font.name = 'Times New Roman'
    run_app.font.size = Pt(16)
    run_app.bold = True

    add_section_h1(doc, "A.1", "User Manual & Operational Guide")
    add_body_p(doc, "SentixAI is designed for intuitive web navigation. Below is the step-by-step operational guide for users and evaluators:")
    add_bullet_item(doc, "Open a web browser and navigate to the public URL: https://kavyasamarth.pythonanywhere.com. Review the system workflow cards explaining text ingestion, preprocessing, NLP classification, and relational archival.", "1. Accessing the Homepage: ")
    add_bullet_item(doc, "Click 'Login' in the navigation bar. Log in using pre-seeded researcher credentials (username: 'nirbhay') or administrator credentials (username: 'admin', password: 'admin123'). Alternatively, register a new account on '/register'.", "2. Authentication: ")
    add_bullet_item(doc, "Navigate to '/analyze'. Enter or paste political statements, social media posts, or legislative remarks into the textarea. Optionally click any of the 3 pre-loaded demonstration cards (Positive, Negative, Neutral) to auto-populate sample text. Click 'Analyze & Record Entry'.", "3. Analyzing Political Text: ")
    add_bullet_item(doc, "Inspect the computed inference diagnostics: the dominant stance badge, the 3-way distribution bars (Positive %, Neutral %, Negative %), and the side-by-side text comparison displaying the exact words cleaned by the regex pipeline.", "4. Reviewing Inference: ")
    add_bullet_item(doc, "Navigate to '/dashboard' to inspect aggregate statistics (Total Ingested, Positive Stance %, Negative Stance %, Neutral Stance %) and the dynamic Chart.js doughnut visualization.", "5. Inspecting Analytics Dashboard: ")
    add_bullet_item(doc, "Navigate to '/results' to audit the full archival table. Use the stance dropdown to filter records or enter search keywords to isolate specific policy topics. Click the red trash icon to delete personal entries.", "6. Archival History & Search: ")
    add_bullet_item(doc, "Log in as 'admin' and click 'Admin Panel'. Review platform-wide statistics, manage registered user accounts, and moderate global posts across the platform.", "7. Administrator Moderation: ")

    add_section_h1(doc, "A.2", "Quick-Start Deployment Command Reference")
    add_body_p(doc, "For local or cloud deployment, execute the following commands in terminal:")
    add_code_block(doc, 
        "# 1. Clone the GitHub repository\n"
        "git clone https://github.com/samarthkavya2-wq/sentiment-analysis-project.git\n"
        "cd sentiment-analysis-project\n\n"
        "# 2. Create and activate Python virtual environment\n"
        "python -m venv venv\n"
        "source venv/bin/activate       # On Linux/PythonAnywhere\n"
        "# venv\\Scripts\\activate       # On Windows PowerShell\n\n"
        "# 3. Install production dependencies\n"
        "pip install -r requirements.txt\n\n"
        "# 4. Initialize database and seed demonstration data\n"
        "python init_db.py\n\n"
        "# 5. Launch local server\n"
        "python app.py"
    , "Command Reference A.1 — Project Setup and Execution Commands")

    add_section_h1(doc, "A.3", "Database Schema Definition (setup.sql)")
    add_body_p(doc, "The complete DDL SQL script utilized to generate the MySQL database schema:")
    add_code_block(doc,
        "-- Database Setup for SentixAI\n"
        "CREATE DATABASE IF NOT EXISTS sentiment_analysis;\n"
        "USE sentiment_analysis;\n\n"
        "CREATE TABLE IF NOT EXISTS users (\n"
        "    id INT AUTO_INCREMENT PRIMARY KEY,\n"
        "    username VARCHAR(50) UNIQUE NOT NULL,\n"
        "    email VARCHAR(100) UNIQUE NOT NULL,\n"
        "    password VARCHAR(255) NOT NULL,\n"
        "    is_admin TINYINT(1) DEFAULT 0,\n"
        "    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP\n"
        ");\n\n"
        "CREATE TABLE IF NOT EXISTS posts (\n"
        "    id INT AUTO_INCREMENT PRIMARY KEY,\n"
        "    user_id INT NOT NULL,\n"
        "    text_content TEXT NOT NULL,\n"
        "    sentiment VARCHAR(20) NOT NULL,\n"
        "    polarity FLOAT NOT NULL,\n"
        "    subjectivity FLOAT NOT NULL,\n"
        "    confidence FLOAT NOT NULL,\n"
        "    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,\n"
        "    FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE\n"
        ");"
    , "Schema Definition A.1 — database/setup.sql")

    print("[+] Post-content constructed.")
