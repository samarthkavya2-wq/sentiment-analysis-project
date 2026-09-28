"""
report_builder/preliminary.py
Generates the academic preliminary pages: Title, Certificate, Declaration,
Acknowledgement, Abstract, Table of Contents, List of Tables, List of Figures.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_builder.common import set_cell_background, set_cell_margins, set_table_borders

def build_preliminary_pages(doc):
    # ---------------------------------------------------------
    # 1. TITLE PAGE
    # ---------------------------------------------------------
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_before = Pt(24)
    p_univ.paragraph_format.space_after = Pt(4)
    r_univ = p_univ.add_run("ACADEMIC PROJECT REPORT\n")
    r_univ.font.name = 'Times New Roman'
    r_univ.font.size = Pt(14)
    r_univ.bold = True
    r_univ.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    r_sub = p_univ.add_run("Submitted in partial fulfillment of the requirements for the Degree of\nBachelor of Science in Computer Science (T.Y.B.Sc. CS)")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    p_proj = doc.add_paragraph()
    p_proj.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_proj.paragraph_format.space_before = Pt(36)
    p_proj.paragraph_format.space_after = Pt(12)
    r_title = p_proj.add_run("A WEB-BASED SYSTEM FOR SENTIMENT ANALYSIS AND FILTERING OF POLITICAL SOCIAL MEDIA TEXT USING PYTHON AND MYSQL\n\n")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x2C, 0x59)

    r_short = p_proj.add_run("System Implementation: SentixAI — Political Sentiment Analysis & Filtering System")
    r_short.font.name = 'Times New Roman'
    r_short.font.size = Pt(13)
    r_short.bold = True
    r_short.font.color.rgb = RGBColor(0x1E, 0x3A, 0x5F)

    p_by = doc.add_paragraph()
    p_by.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_by.paragraph_format.space_before = Pt(48)
    p_by.paragraph_format.space_after = Pt(4)
    r_by = p_by.add_run("Submitted by\n")
    r_by.font.name = 'Times New Roman'
    r_by.font.size = Pt(12)

    r_name = p_by.add_run("KAVYA SAMARTH\n")
    r_name.font.name = 'Times New Roman'
    r_name.font.size = Pt(14)
    r_name.bold = True

    r_roll = p_by.add_run("Roll Number: 9122\nSeat Number: 9122\n")
    r_roll.font.name = 'Times New Roman'
    r_roll.font.size = Pt(12)

    p_guide = doc.add_paragraph()
    p_guide.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_guide.paragraph_format.space_before = Pt(36)
    p_guide.paragraph_format.space_after = Pt(4)
    r_guide = p_guide.add_run("Under the Guidance of\n")
    r_guide.font.name = 'Times New Roman'
    r_guide.font.size = Pt(12)
    r_guide_name = p_guide.add_run("DEPARTMENT OF COMPUTER SCIENCE\n")
    r_guide_name.font.name = 'Times New Roman'
    r_guide_name.font.size = Pt(13)
    r_guide_name.bold = True

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(40)
    p_inst.paragraph_format.space_after = Pt(0)
    r_inst = p_inst.add_run("FACULTY OF SCIENCE & TECHNOLOGY\nACADEMIC YEAR 2026–2027")
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(12)
    r_inst.bold = True

    # ---------------------------------------------------------
    # 2. CERTIFICATE OF APPROVAL
    # ---------------------------------------------------------
    doc.add_page_break()
    p_cert_title = doc.add_paragraph()
    p_cert_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert_title.paragraph_format.space_before = Pt(24)
    p_cert_title.paragraph_format.space_after = Pt(18)
    r_c = p_cert_title.add_run("CERTIFICATE OF APPROVAL")
    r_c.font.name = 'Times New Roman'
    r_c.font.size = Pt(16)
    r_c.bold = True
    r_c.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_cert_b1 = doc.add_paragraph()
    p_cert_b1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cert_b1.paragraph_format.line_spacing = 1.5
    p_cert_b1.paragraph_format.space_after = Pt(12)
    r_cb1 = p_cert_b1.add_run(
        "This is to certify that the project entitled \"A Web-Based System for Sentiment Analysis and Filtering of Political Social Media Text Using Python and MySQL\" (SentixAI) is a bonafide record of independent project work carried out successfully by Kavya Samarth (Roll No: 9122) in partial fulfillment of the requirements for the award of the Degree of Bachelor of Science in Computer Science (T.Y.B.Sc. CS) during the academic year 2026–2027."
    )
    r_cb1.font.name = 'Times New Roman'
    r_cb1.font.size = Pt(12)

    p_cert_b2 = doc.add_paragraph()
    p_cert_b2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cert_b2.paragraph_format.line_spacing = 1.5
    p_cert_b2.paragraph_format.space_after = Pt(40)
    r_cb2 = p_cert_b2.add_run(
        "The project has been examined and evaluated by the internal and external examiners and is hereby approved."
    )
    r_cb2.font.name = 'Times New Roman'
    r_cb2.font.size = Pt(12)

    # Signatures Table
    sig_table = doc.add_table(rows=2, cols=3)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    col_w = Inches(2.1)
    for row in sig_table.rows:
        for c in row.cells:
            c.width = col_w

    sig_labels = [
        ("____________________\nInternal Project Guide", "____________________\nHead of Department", "____________________\nExternal Examiner"),
        ("Date: _____________\nPlace: _____________ ", "College Seal / Stamp", "Date: _____________\nPlace: _____________ ")
    ]
    for r_idx, row_data in enumerate(sig_labels):
        r = sig_table.rows[r_idx]
        for c_idx, text in enumerate(row_data):
            cell = r.cells[c_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(20) if r_idx == 0 else Pt(0)
            p.paragraph_format.line_spacing = 1.2
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)

    # ---------------------------------------------------------
    # 3. CANDIDATE'S DECLARATION
    # ---------------------------------------------------------
    doc.add_page_break()
    p_decl = doc.add_paragraph()
    p_decl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_decl.paragraph_format.space_before = Pt(24)
    p_decl.paragraph_format.space_after = Pt(18)
    r_d = p_decl.add_run("CANDIDATE'S DECLARATION")
    r_d.font.name = 'Times New Roman'
    r_d.font.size = Pt(16)
    r_d.bold = True
    r_d.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_d_body = doc.add_paragraph()
    p_d_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_d_body.paragraph_format.line_spacing = 1.5
    p_d_body.paragraph_format.space_after = Pt(12)
    r_db = p_d_body.add_run(
        "I, Kavya Samarth, student of T.Y.B.Sc. Computer Science (Roll Number: 9122), hereby declare that the project entitled \"A Web-Based System for Sentiment Analysis and Filtering of Political Social Media Text Using Python and MySQL\" (SentixAI) submitted by me to the Department of Computer Science is an original and authentic work done under the supervision and guidance of our project guide.\n\n"
        "I further declare that this work has not formed the basis for the award of any degree, diploma, associateship, fellowship, or other similar title to any candidate in this or any other university or institution. The source code, analytical models, and database designs submitted herewith represent my genuine academic endeavor."
    )
    r_db.font.name = 'Times New Roman'
    r_db.font.size = Pt(12)

    p_d_sig = doc.add_paragraph()
    p_d_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_d_sig.paragraph_format.space_before = Pt(40)
    p_d_sig.paragraph_format.line_spacing = 1.3
    r_ds = p_d_sig.add_run(
        "_________________________\n"
        "Kavya Samarth\n"
        "Roll No: 9122\n"
        "T.Y.B.Sc. (Computer Science)\n"
        "Academic Year 2026–2027\n"
    )
    r_ds.font.name = 'Times New Roman'
    r_ds.font.size = Pt(12)

    # ---------------------------------------------------------
    # 4. ACKNOWLEDGEMENT
    # ---------------------------------------------------------
    doc.add_page_break()
    p_ack = doc.add_paragraph()
    p_ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ack.paragraph_format.space_before = Pt(24)
    p_ack.paragraph_format.space_after = Pt(18)
    r_a = p_ack.add_run("ACKNOWLEDGEMENT")
    r_a.font.name = 'Times New Roman'
    r_a.font.size = Pt(16)
    r_a.bold = True
    r_a.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_a_body = doc.add_paragraph()
    p_a_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_a_body.paragraph_format.line_spacing = 1.5
    p_a_body.paragraph_format.space_after = Pt(12)
    r_ab = p_a_body.add_run(
        "The completion of this academic project report marks an important milestone in my undergraduate studies in Computer Science. I take this opportunity to express my deepest gratitude to all individuals whose constant support, intellectual guidance, and encouragement made this endeavor possible.\n\n"
        "First and foremost, I express my sincere gratitude to our esteemed Principal and the Head of the Department of Computer Science for providing state-of-the-art laboratory infrastructure, computational facilities, and a conducive environment to develop and test full-stack web applications.\n\n"
        "I am profoundly indebted to my Project Guide for their invaluable mentorship, insightful critique, and continuous technical steering throughout the stages of system modeling, Natural Language Processing engine integration, and relational database architecture. Their meticulous review of my design documents helped refine the algorithmic precision of this project.\n\n"
        "I extend my appreciation to all faculty members and technical staff of the Computer Science Department for their constructive feedback during the project progress seminars. Finally, I express heartfelt gratitude to my family and peers for their unwavering encouragement and support throughout this academic year."
    )
    r_ab.font.name = 'Times New Roman'
    r_ab.font.size = Pt(12)

    # ---------------------------------------------------------
    # 5. ABSTRACT
    # ---------------------------------------------------------
    doc.add_page_break()
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs.paragraph_format.space_before = Pt(24)
    p_abs.paragraph_format.space_after = Pt(18)
    r_abs_t = p_abs.add_run("ABSTRACT")
    r_abs_t.font.name = 'Times New Roman'
    r_abs_t.font.size = Pt(16)
    r_abs_t.bold = True
    r_abs_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_abs_body = doc.add_paragraph()
    p_abs_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs_body.paragraph_format.line_spacing = 1.5
    p_abs_body.paragraph_format.space_after = Pt(12)
    r_ab_txt = p_abs_body.add_run(
        "Political discourse on digital social media platforms exerts an unprecedented influence on public policy formulation, electoral democracy, and citizen opinion dynamics. However, the immense volume, unstructured format, and linguistic noise inherent in digital discourse render manual analysis impractical, cost-prohibitive, and vulnerable to subjective human bias. This academic project presents SentixAI, an intelligent web-based system engineered to automatically ingest, sanitize, classify, and filter political sentiment from unstructured textual statements using Python, Natural Language Processing (NLP), and a relational MySQL database.\n\n"
        "The system incorporates a robust Three-Tier Architecture. In the Presentation Layer, dynamic Jinja2 templates, custom responsive CSS, and Chart.js visualizations deliver an intuitive interface for both general researchers and administrative moderators. The Business Logic Layer utilizes an automated Regular Expression (regex) pre-processing pipeline to purge noise such as HTTP/HTTPS hyperlinks, user @mentions, and #hashtags, followed by lexical sentiment polarity and subjectivity scoring via the TextBlob NLP library. To overcome the oversimplification of traditional binary classifiers, SentixAI introduces a fine-grained 3-way distribution algorithm calculating exact integer proportions of Positive, Neutral, and Negative stance summing to 100%.\n\n"
        "The Data Layer comprises a normalized, two-table MySQL schema (`users` and `posts`) featuring strong referential integrity, foreign key cascading constraints, and parameterized SQL queries preventing SQL injection vulnerabilities. Passwords are protected using Werkzeug cryptographic PBKDF2 SHA-256 hashing. The production implementation has been deployed to the PythonAnywhere cloud platform using WSGI process management, establishing an accessible public URL. Comprehensive unit, integration, and security testing confirms high classification reliability, sub-second inference response, and strict cross-user data isolation."
    )
    r_ab_txt.font.name = 'Times New Roman'
    r_ab_txt.font.size = Pt(12)

    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.space_before = Pt(8)
    p_kw.paragraph_format.line_spacing = 1.5
    r_kw_lbl = p_kw.add_run("Keywords: ")
    r_kw_lbl.font.name = 'Times New Roman'
    r_kw_lbl.font.size = Pt(12)
    r_kw_lbl.bold = True
    r_kw = p_kw.add_run("Political Sentiment Analysis, Natural Language Processing, TextBlob, Flask Web Framework, MySQL Relational Database, Regex Preprocessing, Polarity Classification, Cloud WSGI Deployment.")
    r_kw.font.name = 'Times New Roman'
    r_kw.font.size = Pt(12)

    # ---------------------------------------------------------
    # 6. TABLE OF CONTENTS
    # ---------------------------------------------------------
    doc.add_page_break()
    p_toc = doc.add_paragraph()
    p_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc.paragraph_format.space_before = Pt(24)
    p_toc.paragraph_format.space_after = Pt(18)
    r_toc = p_toc.add_run("TABLE OF CONTENTS")
    r_toc.font.name = 'Times New Roman'
    r_toc.font.size = Pt(16)
    r_toc.bold = True
    r_toc.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    toc_items = [
        ("Preliminary Pages (Certificate, Declaration, Acknowledgement, Abstract)", "i - v"),
        ("CHAPTER 1: IDENTIFICATION & FEASIBILITY STUDY", "1"),
        ("    1.1 Introduction / Background", "1"),
        ("    1.2 Problem Identification", "2"),
        ("    1.3 Problem Justification", "3"),
        ("    1.4 Objectives", "4"),
        ("    1.5 Scope & Limitations", "5"),
        ("    1.6 Stakeholder Identification", "6"),
        ("    1.7 Technical Feasibility", "7"),
        ("    1.8 Economic Feasibility", "8"),
        ("    1.9 Operational Feasibility", "9"),
        ("CHAPTER 2: REQUIREMENT ENGINEERING", "10"),
        ("    2.1 Stakeholder Requirements", "10"),
        ("    2.2 Functional Requirements (FR-01 to FR-15)", "11"),
        ("    2.3 Non-Functional Requirements", "14"),
        ("    2.4 Use Case Analysis", "16"),
        ("    2.5 Requirement Prioritization", "17"),
        ("    2.6 System Constraints", "18"),
        ("    2.7 Project Assumptions", "19"),
        ("CHAPTER 3: SOFTWARE DEVELOPMENT LIFE CYCLE", "20"),
        ("    3.1 SDLC Methodology (Waterfall Model)", "20"),
        ("    3.2 Project Phases and Timeline", "22"),
        ("    3.3 Gantt Chart Representation", "24"),
        ("CHAPTER 4: SYSTEM MODELING USING UML", "26"),
        ("    4.1 Use Case Diagram", "26"),
        ("    4.2 Class Diagram", "29"),
        ("    4.3 Sequence Diagrams (Authentication, Ingestion, Search, Moderation)", "31"),
        ("    4.4 Activity Diagram", "36"),
        ("    4.5 Entity-Relationship (ER) Diagram", "38"),
        ("    4.6 Deployment Diagram", "40"),
        ("CHAPTER 5: ARCHITECTURE DESIGN", "42"),
        ("    5.1 Three-Tier Architecture Overview", "42"),
        ("    5.2 Front-End Architecture & Jinja2 Templates", "44"),
        ("    5.3 Backend Architecture & Request-Response Cycle", "46"),
        ("    5.4 Natural Language Processing Pipeline Design", "48"),
        ("    5.5 Database Architecture & Schema Specification", "51"),
        ("    5.6 Flask Route & Endpoint Structure (11 Routes)", "54"),
        ("    5.7 Security Architecture", "56"),
        ("CHAPTER 6: APPLICATION DEVELOPMENT", "58"),
        ("    6.1 Front-End Implementation & Design System", "58"),
        ("    6.2 Backend Implementation Logic", "60"),
        ("    6.3 Database Integration & Relational Queries", "62"),
        ("    6.4 Input Validation & Error Handling", "64"),
        ("    6.5 Module-Wise Implementation Breakdown", "66"),
        ("    6.6 Visual Representations of User Interfaces", "70"),
        ("    6.7 Annotated Source Code Excerpts", "76"),
        ("CHAPTER 7: INTEGRATION AND SYSTEM TESTING", "82"),
        ("    7.1 Testing Strategy & Methodology", "82"),
        ("    7.2 Unit Testing", "83"),
        ("    7.3 Black-Box Testing", "84"),
        ("    7.4 Integration Testing", "85"),
        ("    7.5 System Testing", "86"),
        ("    7.6 Test Case Execution Matrix (Test Cases TC-01 to TC-18)", "87"),
        ("    7.7 Defect Tracking & Bug Resolution History", "93"),
        ("CHAPTER 8: DEPLOYMENT", "95"),
        ("    8.1 Deployment Environment & Topology", "95"),
        ("    8.2 Production Configuration Hardening", "96"),
        ("    8.3 Cloud Database Schema Migration & Seeding", "97"),
        ("    8.4 Virtual Environment & Dependency Isolation", "98"),
        ("    8.5 Step-by-Step PythonAnywhere Hosting Procedure", "99"),
        ("    8.6 Version Control & GitHub Repository Integration", "102"),
        ("    8.7 Public Access & Live System Verification", "103"),
        ("CHAPTER 9: PERFORMANCE AND SECURITY TESTING", "104"),
        ("    9.1 Input Validation & SQL Injection Verification", "104"),
        ("    9.2 Authentication & Hash Verification Testing", "105"),
        ("    9.3 Role-Based Authorization & Privilege Escalation Testing", "106"),
        ("    9.4 Cross-User Data Tampering Prevention Testing", "107"),
        ("    9.5 System Response & Inference Latency Benchmark", "108"),
        ("    9.6 Known Security & Technical Limitations", "109"),
        ("CHAPTER 10: RESULT AND DISCUSSION", "110"),
        ("    10.1 Test Execution Summary", "110"),
        ("    10.2 Functional Evaluation & Key Achievements", "111"),
        ("    10.3 Comparison of Cleaned vs Uncleaned Text Polarity", "112"),
        ("    10.4 System Evaluation & Review", "114"),
        ("    10.5 Academic Limitations", "115"),
        ("    10.6 Future Scope & Enhancements", "116"),
        ("    10.7 Conclusion", "117"),
        ("REFERENCES", "118"),
        ("APPENDIX: USER MANUAL & DEPLOYMENT REFERENCE", "120")
    ]

    t_toc = doc.add_table(rows=len(toc_items), cols=2)
    t_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_toc.autofit = False
    for r_idx, (title, page) in enumerate(toc_items):
        row = t_toc.rows[r_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(5.4)
        c1.width = Inches(0.9)
        set_cell_margins(c0, 20, 20, 40, 40)
        set_cell_margins(c1, 20, 20, 40, 40)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(title)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(10.5)
        if "CHAPTER" in title:
            r0.bold = True

        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(page)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)
        if "CHAPTER" in title:
            r1.bold = True

    # ---------------------------------------------------------
    # 7. LIST OF TABLES
    # ---------------------------------------------------------
    doc.add_page_break()
    p_lot = doc.add_paragraph()
    p_lot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lot.paragraph_format.space_before = Pt(24)
    p_lot.paragraph_format.space_after = Pt(18)
    r_lot = p_lot.add_run("LIST OF TABLES")
    r_lot.font.name = 'Times New Roman'
    r_lot.font.size = Pt(16)
    r_lot.bold = True
    r_lot.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    tables_list = [
        ("Table 1.1", "Economic Feasibility and Open-Source Cost Analysis"),
        ("Table 2.1", "Functional Requirements Specification (FR-01 to FR-15)"),
        ("Table 2.2", "Non-Functional Requirements Specification"),
        ("Table 2.3", "Requirement Prioritization Matrix (MoSCoW)"),
        ("Table 3.1", "Waterfall Project Lifecycle Timeline and Milestones"),
        ("Table 5.1", "MySQL Database Schema: 'users' Table Definition"),
        ("Table 5.2", "MySQL Database Schema: 'posts' Table Definition"),
        ("Table 5.3", "Complete Flask Route and Endpoint Specification (11 Routes)"),
        ("Table 6.1", "Text Preprocessing Regex Cleaning Transformations"),
        ("Table 6.2", "Academic Demonstration Accounts Seeded in Database"),
        ("Table 7.1", "Comprehensive Test Case Execution Matrix (TC-01 to TC-18)"),
        ("Table 7.2", "Resolved Software Defects and Bug Tracking History"),
        ("Table 8.1", "Hosting Platforms Comparative Evaluation"),
        ("Table 9.1", "Inference and Response Latency Benchmarks"),
        ("Table 10.1", "Experimental Impact of Regex Cleaning on Sentiment Scoring")
    ]

    t_lot = doc.add_table(rows=len(tables_list), cols=2)
    t_lot.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_lot.autofit = False
    for r_idx, (num, desc) in enumerate(tables_list):
        row = t_lot.rows[r_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.5)
        c1.width = Inches(4.8)
        set_cell_margins(c0, 30, 30, 40, 40)
        set_cell_margins(c1, 30, 30, 40, 40)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(num)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(10.5)
        r0.bold = True

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(desc)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)

    # ---------------------------------------------------------
    # 8. LIST OF FIGURES
    # ---------------------------------------------------------
    doc.add_page_break()
    p_lof = doc.add_paragraph()
    p_lof.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lof.paragraph_format.space_before = Pt(24)
    p_lof.paragraph_format.space_after = Pt(18)
    r_lof = p_lof.add_run("LIST OF FIGURES")
    r_lof.font.name = 'Times New Roman'
    r_lof.font.size = Pt(16)
    r_lof.bold = True
    r_lof.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    figures_list = [
        ("Figure 3.1", "Waterfall Software Development Life Cycle Phases"),
        ("Figure 3.2", "Project Gantt Chart Timeline (Semester-V 2026-27)"),
        ("Figure 4.1", "Use Case Diagram for SentixAI"),
        ("Figure 4.2", "System Class Diagram"),
        ("Figure 4.3a", "Sequence Diagram: User Registration and Authentication Flow"),
        ("Figure 4.3b", "Sequence Diagram: Political Text Ingestion and Sentiment Analysis"),
        ("Figure 4.3c", "Sequence Diagram: Multi-Dimensional History Search & Filtering"),
        ("Figure 4.3d", "Sequence Diagram: Administrator Moderation and Account Deletion"),
        ("Figure 4.4", "End-to-End System Activity Diagram"),
        ("Figure 4.5", "Entity-Relationship (ER) Diagram (users 1:M posts)"),
        ("Figure 4.6", "Physical Deployment Diagram (3-Tier Topology)"),
        ("Figure 5.1", "Three-Tier Architectural Decomposition"),
        ("Figure 5.2", "Natural Language Processing (NLP) Pipeline Architecture"),
        ("Figure 6.1", "User Interface Wireframe: Public Landing Page (index.html)"),
        ("Figure 6.2", "User Interface Wireframe: User Login Page (login.html)"),
        ("Figure 6.3", "User Interface Wireframe: Registration Page (register.html)"),
        ("Figure 6.4", "User Interface Wireframe: Text Sentiment Classification (analyze.html)"),
        ("Figure 6.5", "User Interface Wireframe: Diagnostic Results with 3-Way Distribution"),
        ("Figure 6.6", "User Interface Wireframe: Sentiment Intelligence Dashboard (dashboard.html)"),
        ("Figure 6.7", "User Interface Wireframe: Analysis History & Live Filtering (results.html)"),
        ("Figure 6.8", "User Interface Wireframe: Administrator & Moderation Panel (admin.html)")
    ]

    t_lof = doc.add_table(rows=len(figures_list), cols=2)
    t_lof.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_lof.autofit = False
    for r_idx, (num, desc) in enumerate(figures_list):
        row = t_lof.rows[r_idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.5)
        c1.width = Inches(4.8)
        set_cell_margins(c0, 30, 30, 40, 40)
        set_cell_margins(c1, 30, 30, 40, 40)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(num)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(10.5)
        r0.bold = True

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(desc)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)

    print("[+] Preliminary pages successfully constructed.")
